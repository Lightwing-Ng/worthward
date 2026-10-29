"""Anonymous regressions for authoritative Schwab execution reconciliation.

Code version: v1.0.0
"""

from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
import pytest

from app.services.investment.importing.brokers.zircon_hk import (
    build_zircon_hk_template_xlsx,
)
from app.services.investment.investment_record_basics import (
    build_normalized_transaction_view,
)
from app.services.investment.importing.merge.schwab_provisional import (
    reconcile_schwab_provisional_trades,
)
from app.services.investment_import import (
    merge_investment_payloads,
    parse_investment_payload,
)
from tests.python.services.test_schwab_provisional_reconciliation import (
    buy_csv_payload,
    dram_buys,
    paired_csv_payload,
    provisional_buy_payload,
    provisional_payload,
)


TRADE_DAY = "10/13/2026"


def _set_trade_economics(
    transaction: dict,
    price: str,
    *,
    quantity: str | None = None,
    fee: str = "0",
) -> None:
    """Adjust existing anonymous fixtures without creating another record factory."""
    side = transaction["type"]
    units = abs(Decimal(quantity or transaction["quantity_raw"]))
    signed_units = -units if side == "sell" else units
    unit_price = Decimal(price)
    commission = -abs(Decimal(fee))
    gross = units * unit_price * (1 if side == "sell" else -1)
    net = gross + commission
    transaction.update(
        quantity_raw=str(signed_units),
        quantity_abs=str(units),
        price_raw=price,
        commission_raw=str(commission),
        gross_amount_raw=str(gross),
        net_amount_raw=str(net),
        normalized=build_normalized_transaction_view(
            side,
            signed_units,
            unit_price,
            gross,
            commission if commission else None,
            net,
            is_cash_flow_override=False,
            side_override=side,
        ),
    )


def _limit_buy(price: str = "61.00", *, sequence: int = 1) -> dict:
    payload = provisional_buy_payload()
    transaction = payload["transactions"][0]
    _set_trade_economics(transaction, price)
    source = transaction["source"]
    offset = timedelta(minutes=sequence - 1)
    entered = datetime.fromisoformat(source["source_datetime_raw"]) + offset
    transaction["datetime"] = (
        datetime.fromisoformat(transaction["datetime"]) + offset
    ).strftime("%Y-%m-%d %H:%M:%S")
    source.update(
        source_datetime_raw=entered.strftime("%Y-%m-%d %H:%M:%S"),
        reference_id=f"TOS-PAGE-{entered:%Y%m%d-%H%M%S}-DRAM-BUY-{sequence}",
        source_row=sequence + 1,
        order_type="limit",
        limit_price_raw=price,
        execution_status="FILLED",
    )
    return payload


def _limit_sale() -> dict:
    payload = provisional_payload()
    payload["transactions"][0]["source"].update(
        order_type="limit",
        limit_price_raw="51.00",
        execution_status="FILLED",
    )
    return payload


def _day_only_buy(*, file_kind: str = "manual_investment_xlsx") -> dict:
    payload = provisional_buy_payload(
        reference="USER-REPORTED-20261013-DRAM-BUY-6000-5-1"
    )
    transaction = payload["transactions"][0]
    _set_trade_economics(transaction, "60.00")
    transaction.update(date="2026-10-13", datetime="2026-10-13 11:00:00")
    transaction["source"].update(
        file_kind=file_kind,
        source_datetime_raw="2026-10-13 23:00:00",
        source_time_defaulted_to_2300=True,
        source_has_intraday_timestamp=False,
        datetime_precision="day",
        execution_evidence_source="user_confirmed_trade_report",
        requires_manual_official_reconciliation=True,
    )
    return payload


def _official_buys(*prices: str, fee: str = "0") -> dict:
    return buy_csv_payload(buys=tuple((TRADE_DAY, price, fee) for price in prices))


def _linked_rows(payload: dict) -> dict[str, dict]:
    return {
        row["source"]["superseded_manual_reference_id"]: row
        for row in payload["transactions"]
        if (row.get("source") or {}).get("superseded_manual_reference_id")
    }


def test_unique_buy_price_improvement_preserves_official_economics_and_evidence():
    manual = _limit_buy()
    official = _official_buys("60.98", fee="0.07")
    manual_before = deepcopy(manual)
    official_before = deepcopy(official)

    merged = merge_investment_payloads(manual, official)

    assert len(dram_buys(merged)) == 1
    row = dram_buys(merged)[0]
    original = dram_buys(official_before)[0]
    for field in (
        "price_raw",
        "commission_raw",
        "gross_amount_raw",
        "net_amount_raw",
        "normalized",
    ):
        assert row[field] == original[field]
    assert row["source"]["file_kind"] == "schwab_csv"
    if "source_file_sha256" in original["source"]:
        assert (
            row["source"]["source_file_sha256"]
            == original["source"]["source_file_sha256"]
        )
    assert row["source"]["schwab_provisional_match_method"] == "limit_price_improvement"
    assert (
        row["source"]["superseded_manual_reference_id"]
        == manual["transactions"][0]["source"]["reference_id"]
    )
    assert (
        row["source"]["superseded_manual_source_sha256"]
        == manual["source_artifacts"][0]["sha256"]
    )
    assert {item["sha256"] for item in merged["source_artifacts"]} == {
        item["sha256"]
        for payload in (manual, official)
        for item in payload["source_artifacts"]
    }
    archived = row["source"]["superseded_manual_trade"]
    assert archived["price_raw"] == manual_before["transactions"][0]["price_raw"]
    assert (
        archived["source"]["reference_id"]
        == manual_before["transactions"][0]["source"]["reference_id"]
    )


@pytest.mark.parametrize("replay_manual_first", [False, True])
def test_improved_canonical_execution_accepts_its_original_manual_replay(
    replay_manual_first,
):
    manual = _limit_buy()
    official = _official_buys("60.98", fee="0.07")
    replayed = merge_investment_payloads(manual, official)
    replay_order = (manual, official) if replay_manual_first else (official, manual)
    for incoming in replay_order:
        replayed = merge_investment_payloads(replayed, incoming)

    assert len(dram_buys(replayed)) == 1
    assert dram_buys(replayed)[0]["price_raw"] == "60.98"
    assert dram_buys(replayed)[0]["commission_raw"] == "-0.07"


@pytest.mark.parametrize("new_limit", ["60.98", "62.00"])
def test_another_manual_reference_cannot_claim_a_previously_linked_execution(
    new_limit,
):
    original_manual = _limit_buy()
    linked = merge_investment_payloads(original_manual, _official_buys("60.98"))
    unrelated_manual = _limit_buy(new_limit, sequence=2)

    with pytest.raises(ValueError, match="provisional.*Schwab"):
        merge_investment_payloads(linked, unrelated_manual)


def test_unique_sell_price_improvement_uses_official_net_proceeds():
    manual = _limit_sale()
    official = paired_csv_payload(
        include_sale=True, sale_price="51.04", sale_fee="0.03"
    )

    merged = merge_investment_payloads(manual, official)

    sales = [row for row in merged["transactions"] if row["type"] == "sell"]
    assert len(sales) == 1
    assert Decimal(sales[0]["net_amount_raw"]) == Decimal("51.01")
    assert sales[0]["commission_raw"] == "-0.03"
    assert (
        sales[0]["source"]["schwab_provisional_match_method"]
        == "limit_price_improvement"
    )


@pytest.mark.parametrize("side", ["buy", "sell"])
def test_worse_than_limit_execution_stays_blocked(side):
    manual = _limit_buy() if side == "buy" else _limit_sale()
    official = (
        _official_buys("61.01")
        if side == "buy"
        else paired_csv_payload(include_sale=True, sale_price="50.99")
    )

    with pytest.raises(ValueError, match="provisional.*Schwab"):
        merge_investment_payloads(manual, official)


@pytest.mark.parametrize(
    "missing_field", ["order_type", "limit_price_raw", "execution_status"]
)
def test_improvement_requires_complete_structured_limit_evidence(missing_field):
    manual = _limit_buy()
    manual["transactions"][0]["source"].pop(missing_field)

    with pytest.raises(ValueError, match="provisional.*Schwab"):
        merge_investment_payloads(manual, _official_buys("60.98"))


def test_stated_limit_must_equal_the_provisional_price():
    manual = _limit_buy()
    manual["transactions"][0]["source"]["limit_price_raw"] = "61.01"

    with pytest.raises(ValueError, match="provisional.*Schwab"):
        merge_investment_payloads(manual, _official_buys("60.98"))


def test_limit_evidence_normalizes_order_type_and_fill_status():
    manual = _limit_buy()
    manual["transactions"][0]["source"].update(
        order_type="LIMIT", execution_status="filled"
    )

    merged = merge_investment_payloads(manual, _official_buys("60.98"))

    assert len(dram_buys(merged)) == 1


@pytest.mark.parametrize("reverse_manuals", [False, True])
@pytest.mark.parametrize("reverse_csv_rows", [False, True])
def test_global_exact_reservations_precede_limit_improvement(
    reverse_manuals, reverse_csv_rows
):
    limit = _limit_buy()
    exact = _limit_buy("60.00", sequence=2)
    manual = deepcopy(limit)
    manual["transactions"] += exact["transactions"]
    if reverse_manuals:
        manual["transactions"].reverse()
    prices = ("60.00", "60.98") if reverse_csv_rows else ("60.98", "60.00")

    merged = merge_investment_payloads(manual, _official_buys(*prices))

    linked = _linked_rows(merged)
    assert len(dram_buys(merged)) == 2
    assert (
        linked[limit["transactions"][0]["source"]["reference_id"]]["price_raw"]
        == "60.98"
    )
    assert (
        linked[exact["transactions"][0]["source"]["reference_id"]]["price_raw"]
        == "60.00"
    )


def test_exact_day_only_fill_is_reserved_before_a_limit_order_can_claim_it():
    manual = _limit_buy()
    day_only = _day_only_buy()
    manual["transactions"] += day_only["transactions"]

    merged = merge_investment_payloads(manual, _official_buys("60.00", "60.98"))

    linked = _linked_rows(merged)
    assert len(dram_buys(merged)) == 2
    assert (
        linked[day_only["transactions"][0]["source"]["reference_id"]]["price_raw"]
        == "60.00"
    )


def test_distinct_extra_fill_is_retained_before_and_after_reimport():
    manual = _limit_buy()
    first = _official_buys("60.98")
    later = _official_buys("60.98", "61.50")

    merged = merge_investment_payloads(manual, later)
    assert sorted(Decimal(row["price_raw"]) for row in dram_buys(merged)) == [
        Decimal("60.98"),
        Decimal("61.50"),
    ]
    assert len(_linked_rows(merged)) == 1
    replayed = merge_investment_payloads(merged, later)
    assert len(dram_buys(replayed)) == 2

    initially_linked = merge_investment_payloads(manual, first)
    extended = merge_investment_payloads(initially_linked, later)
    assert len(dram_buys(extended)) == 2


@pytest.mark.parametrize("prices", [("60.98", "60.98"), ("60.98", "60.50")])
def test_multiple_improved_candidates_are_ambiguous_instead_of_choosing_nearest(prices):
    with pytest.raises(ValueError, match="provisional.*Schwab"):
        merge_investment_payloads(_limit_buy(), _official_buys(*prices))


@pytest.mark.parametrize(
    "file_kind", ["manual_investment_xlsx", "zircon_hk_manual_xlsx"]
)
@pytest.mark.parametrize("precise_official_time", [False, True])
def test_day_only_user_report_reconciles_without_fabricating_execution_seconds(
    file_kind, precise_official_time
):
    manual = _day_only_buy(file_kind=file_kind)
    official = _official_buys("60.00", fee="0.07")
    if precise_official_time:
        trade = dram_buys(official)[0]
        trade["datetime"] = "2026-10-13 09:14:32"
        trade["source"].update(
            source_has_intraday_timestamp=True, datetime_precision="second"
        )

    merged = merge_investment_payloads(manual, official)

    assert len(dram_buys(merged)) == 1
    row = dram_buys(merged)[0]
    source = row["source"]
    original_source = manual["transactions"][0]["source"]
    assert source["file_kind"] == "schwab_csv"
    assert source["superseded_manual_reference_id"] == original_source["reference_id"]
    assert not source["superseded_manual_reference_id"].startswith("TOS-PAGE-")
    assert source["superseded_manual_source_datetime_raw"] == "2026-10-13 23:00:00"
    assert source["superseded_manual_source_timezone"] == "Asia/Hong_Kong"
    assert source["superseded_manual_datetime_precision"] == "day"
    assert set(source["superseded_manual_eligible_dates"]) == {"2026-10-13"}
    assert source["schwab_provisional_match_method"] == "exact"
    assert row["datetime"] == dram_buys(official)[0]["datetime"]
    assert Decimal(row["net_amount_raw"]) == Decimal("-300.07")
    assert len(dram_buys(merge_investment_payloads(merged, official))) == 1
    assert len(dram_buys(merge_investment_payloads(merged, manual))) == 1


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("source_has_intraday_timestamp", True),
        ("datetime_precision", "second"),
        ("execution_evidence_source", "unverified"),
        ("requires_manual_official_reconciliation", False),
        ("reference_id", "unrelated-manual-reference"),
    ],
)
def test_unverified_day_only_metadata_does_not_authorize_reconciliation(field, value):
    manual = _day_only_buy()
    manual["transactions"][0]["source"][field] = value

    existing, _, removed = reconcile_schwab_provisional_trades(
        manual["transactions"], _official_buys("60.00")["transactions"]
    )

    assert removed == 0
    assert len(existing) == 1


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("source_time_defaulted_to_2300", True),
        ("source_has_intraday_timestamp", False),
    ],
)
def test_tos_reference_cannot_promote_a_default_time_into_execution_evidence(
    field, value
):
    manual = _limit_buy()
    manual["transactions"][0]["source"][field] = value
    official = _official_buys("61.00")

    existing, _, removed = reconcile_schwab_provisional_trades(
        manual["transactions"], official["transactions"]
    )

    assert removed == 0
    assert existing == manual["transactions"]
    with pytest.raises(ValueError, match="Schwab merged holdings"):
        merge_investment_payloads(manual, official)


def test_day_only_user_report_requires_exact_price_and_unambiguous_quantity():
    for official in (_official_buys("59.98"), _official_buys("60.00", "60.00")):
        with pytest.raises(ValueError, match="provisional.*Schwab"):
            merge_investment_payloads(_day_only_buy(), official)


@pytest.mark.parametrize("manual_factory", [_limit_buy, _day_only_buy])
@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("account", "Individual ...999"),
        ("ticker", "OTHER"),
        ("currency", "HKD"),
        ("type", "sell"),
    ],
)
def test_manual_trades_never_claim_another_account_security_currency_or_side(
    manual_factory, field, value
):
    manual = manual_factory()
    price = manual["transactions"][0]["price_raw"]
    official = _official_buys(price)
    trade = dram_buys(official)[0]
    trade[field] = value
    if field == "account":
        trade["source"]["account"] = value

    existing, _, removed = reconcile_schwab_provisional_trades(
        manual["transactions"], official["transactions"]
    )

    assert removed == 0
    assert len(existing) == 1


@pytest.mark.parametrize("manual_factory", [_limit_buy, _day_only_buy])
def test_different_quantity_cannot_be_absorbed_as_a_matching_fill(manual_factory):
    manual = manual_factory()
    trade = manual["transactions"][0]
    _set_trade_economics(trade, trade["price_raw"], quantity="6")

    with pytest.raises(ValueError, match="provisional.*Schwab"):
        reconcile_schwab_provisional_trades(
            manual["transactions"], _official_buys(trade["price_raw"])["transactions"]
        )


def test_two_manual_orders_cannot_share_one_official_execution():
    manual = _limit_buy()
    manual["transactions"] += _limit_buy(sequence=2)["transactions"]

    with pytest.raises(ValueError, match="provisional.*Schwab"):
        merge_investment_payloads(manual, _official_buys("60.98"))


@pytest.mark.parametrize(
    "manual_factory,price", [(_limit_buy, "60.98"), (_day_only_buy, "60.00")]
)
@pytest.mark.parametrize("revision", ["fee", "price", "quantity", "missing"])
def test_linked_official_execution_cannot_be_revised_or_replaced(
    manual_factory, price, revision
):
    official = _official_buys(price, fee="0.07")
    linked = merge_investment_payloads(manual_factory(), official)
    if revision == "missing":
        revised = _official_buys("61.50", fee="0.07")
    elif revision == "price":
        revised = _official_buys(str(Decimal(price) - Decimal("0.01")), fee="0.07")
    elif revision == "fee":
        revised = _official_buys(price, fee="0.08")
    else:
        revised = deepcopy(official)
        _set_trade_economics(dram_buys(revised)[0], price, quantity="6", fee="0.07")

    with pytest.raises(ValueError, match="Schwab.*conflicting|Schwab.*missing"):
        merge_investment_payloads(linked, revised)


@pytest.mark.parametrize(
    ("linked_day", "revised_day"),
    [("10/12/2026", "10/13/2026"), ("10/13/2026", "10/12/2026")],
)
def test_day_only_link_preserves_its_original_adjacent_date_window(
    linked_day, revised_day
):
    manual = _day_only_buy()
    # A legacy day-precision record can already have adjacent ledger and source dates.
    manual["transactions"][0].update(date="2026-10-12", datetime="2026-10-12 20:00:00")
    official = buy_csv_payload(buys=((linked_day, "60.00", "0.07"),))
    linked = merge_investment_payloads(manual, official)
    row = dram_buys(linked)[0]
    assert set(row["source"]["superseded_manual_eligible_dates"]) == {
        "2026-10-12",
        "2026-10-13",
    }
    revised = buy_csv_payload(buys=((revised_day, "60.00", "0.07"),))

    with pytest.raises(ValueError, match="Schwab.*conflicting|Schwab.*missing"):
        merge_investment_payloads(linked, revised)


def test_unverified_archived_reference_and_saved_date_list_cannot_expand_a_link():
    manual = _day_only_buy()
    linked = merge_investment_payloads(manual, _official_buys("60.00"))
    row = dram_buys(linked)[0]
    row["source"]["superseded_manual_eligible_dates"] = [
        "2026-10-12",
        "2026-10-13",
    ]
    archived = row["source"]["superseded_manual_trade"]
    archived["date"] = "2026-10-12"
    archived["source"]["reference_id"] = "USER-REPORTED-UNRELATED-REFERENCE"
    different_day = buy_csv_payload(buys=(("10/12/2026", "60.00", "0"),))

    existing, incoming, removed = reconcile_schwab_provisional_trades(
        [row], different_day["transactions"]
    )

    assert removed == 0
    assert existing[0]["date"] == "2026-10-13"
    assert len([item for item in incoming if item.get("ticker") == "DRAM"]) == 1


def _real_day_only_workbook():
    workbook = load_workbook(BytesIO(build_zircon_hk_template_xlsx()))
    sheet = workbook["Transactions"]
    reference = "USER-REPORTED-20261013-DRAM-BUY-6000-5-1"
    values = [
        "schwab",
        "Individual ...001",
        date(2026, 10, 13),
        "Buy",
        "USD",
        "DRAM",
        5,
        60,
        None,
        0,
        "Anonymous date-only purchase",
        reference,
    ]
    for column, value in enumerate(values, 1):
        sheet.cell(row=2, column=column, value=value)
    sheet.cell(row=2, column=3).number_format = "yyyy-mm-dd"
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


def _parse_workbook(content):
    return parse_investment_payload(
        "zircon_hk",
        "manual_xlsx",
        xlsx_bytes=content,
        filename="anonymous-date-only.xlsx",
    )


def _linked_real_workbook(*, precise_official=False):
    content = _real_day_only_workbook()
    confirmed = _parse_workbook(content)
    confirmed["transactions"][0]["source"].update(
        source_has_intraday_timestamp=False,
        datetime_precision="day",
        execution_evidence_source="user_confirmed_trade_report",
        requires_manual_official_reconciliation=True,
    )
    confirmed["transactions"][0]["commission_raw"] = "0"
    confirmed["transactions"][0]["normalized"].update(
        commission="0",
        commission_display="0",
    )
    official = _official_buys("60.00")
    if precise_official:
        trade = dram_buys(official)[0]
        trade["datetime"] = "2026-10-13 08:14:32"
        trade["source"].update(
            source_has_intraday_timestamp=True,
            datetime_precision="second",
        )
    return content, merge_investment_payloads(confirmed, official)


@pytest.mark.parametrize("manual_first", [False, True])
@pytest.mark.parametrize("precise_official", [False, True])
def test_original_date_only_workbook_reparse_reuses_verified_consumed_reference(
    manual_first,
    precise_official,
):
    content, linked = _linked_real_workbook(precise_official=precise_official)
    raw_reparse = _parse_workbook(content)
    source = raw_reparse["transactions"][0]["source"]
    assert "source_has_intraday_timestamp" not in source
    assert "datetime_precision" not in source
    assert "execution_evidence_source" not in source
    assert "requires_manual_official_reconciliation" not in source
    assert "commission_raw" not in raw_reparse["transactions"][0]
    expected = deepcopy(linked["transactions"])
    left, right = (raw_reparse, linked) if manual_first else (linked, raw_reparse)

    merged = merge_investment_payloads(left, right)

    assert len(dram_buys(merged)) == 1
    assert merged["transactions"] == expected
    assert dram_buys(merged)[0]["source"]["file_kind"] == "schwab_csv"
    repeated = merge_investment_payloads(merged, _parse_workbook(content))
    assert repeated["transactions"] == expected


@pytest.mark.parametrize(
    ("location", "field", "value"),
    [
        ("source", "source_file_sha256", "f" * 64),
        ("source", "source_file_sha256", "not-a-digest"),
        ("source", "source_sheet", "Changed"),
        ("source", "source_row", 3),
        ("source", "source_row", 0),
        ("source", "source_row", True),
        ("source", "reference_id", "USER-REPORTED-CHANGED-REFERENCE"),
        ("source", "source_datetime_raw", "2026-10-13 22:00:00"),
        ("source", "source_timezone", "UTC"),
        ("source", "source_has_intraday_timestamp", True),
        ("source", "datetime_precision", "second"),
        ("source", "execution_evidence_source", "unverified"),
        ("source", "requires_manual_official_reconciliation", False),
        ("source", "source_time_defaulted_to_2300", False),
        ("record", "account", "Individual ...002"),
        ("record", "ticker", "OTHER"),
        ("record", "currency", "HKD"),
        ("record", "type", "sell"),
        ("record", "type", "dividend"),
        ("record", "broker", "ibkr"),
        ("record", "quantity_raw", "6"),
        ("record", "price_raw", "60.01"),
        ("record", "gross_amount_raw", "-299.99"),
        ("record", "net_amount_raw", "-300.01"),
        ("record", "commission_raw", "-0.01"),
        ("normalized", "net_amount", "-300.01"),
        ("normalized", "commission", "0.01"),
        ("normalized", "display_quantity", "6"),
    ],
)
def test_changed_consumed_workbook_identity_or_economics_fails_review(
    location,
    field,
    value,
):
    content, linked = _linked_real_workbook()
    raw_reparse = _parse_workbook(content)
    record = raw_reparse["transactions"][0]
    target = record if location == "record" else record[location]
    target[field] = value

    with pytest.raises(ValueError, match="consumed provisional Schwab reference"):
        merge_investment_payloads(linked, raw_reparse)


def test_new_workbook_and_new_reference_remain_a_distinct_manual_input():
    content, linked = _linked_real_workbook()
    distinct = _parse_workbook(content)
    source = distinct["transactions"][0]["source"]
    source["source_file_sha256"] = "f" * 64
    source["reference_id"] = "USER-REPORTED-DISTINCT-INPUT"

    _existing, retained, removed = reconcile_schwab_provisional_trades(
        linked["transactions"], distinct["transactions"]
    )

    assert removed == 0
    assert retained == distinct["transactions"]


def test_same_reference_on_distinct_other_broker_workbook_is_not_claimed():
    content, linked = _linked_real_workbook()
    distinct = _parse_workbook(content)
    row = distinct["transactions"][0]
    row.update(broker="ibkr", account="U00000001")
    row["source"].update(
        source_file_sha256="f" * 64,
        broker="ibkr",
        account="U00000001",
    )

    _existing, retained, removed = reconcile_schwab_provisional_trades(
        linked["transactions"], distinct["transactions"]
    )

    assert removed == 0
    assert retained == distinct["transactions"]


@pytest.mark.parametrize("invalid_archive", [False, True])
def test_unverified_consumed_user_report_cannot_recover_missing_parser_enrichment(
    invalid_archive,
):
    content, linked = _linked_real_workbook()
    source = dram_buys(linked)[0]["source"]
    if invalid_archive:
        source["superseded_manual_trade"]["source"]["execution_evidence_source"] = (
            "unverified"
        )
    else:
        source.pop("superseded_manual_trade")
    with pytest.raises(ValueError, match="consumed provisional Schwab reference"):
        merge_investment_payloads(linked, _parse_workbook(content))


def test_legacy_recognized_tos_link_without_full_archive_keeps_exact_replay():
    manual = _limit_buy()
    official = _official_buys("61.00", fee="0.07")
    linked = merge_investment_payloads(manual, official)
    dram_buys(linked)[0]["source"].pop("superseded_manual_trade")

    replayed = merge_investment_payloads(linked, manual)

    assert len(dram_buys(replayed)) == 1
    assert dram_buys(replayed)[0]["price_raw"] == "61.00"
    assert dram_buys(replayed)[0]["commission_raw"] == "-0.07"
    with pytest.raises(ValueError, match="Schwab.*conflicting"):
        merge_investment_payloads(replayed, _official_buys("61.00", fee="0.08"))
