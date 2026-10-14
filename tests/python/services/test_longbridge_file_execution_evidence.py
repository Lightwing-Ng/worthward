"""Longbridge paired-file quantity, cash, and booking-time evidence.

Code version: v0.1.0
"""

from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from app.services.investment.importing import basics, compat
from app.services.investment.importing.brokers.longbridge_files import (
    build_investment_payload_from_longbridge_sg_files,
    build_investment_payload_from_longbridge_hk_files,
)
from tests.python.services.test_longbridge_sg_import import (
    _build_history_orders_workbook_bytes,
)


def _orders(
    *,
    quantity=10,
    price=10,
    time="2024-11-15 10:00:00 ET",
    status="Filled",
    empty=False,
):
    workbook = load_workbook(BytesIO(_build_history_orders_workbook_bytes()))
    sheet = workbook["Orders"]
    headers = [c.value for c in sheet[1]]
    sheet.delete_rows(2, sheet.max_row)
    for name in list(workbook.sheetnames):
        if name != "Orders":
            del workbook[name]
    if not empty:
        values = {
            "Order Status": status,
            "Market": "US",
            "Symbol": "TQQQ",
            "Stock Name": "Fund",
            "Direction": "Buy",
            "Order Type": "Limit",
            "Quantity": quantity,
            "Price": price,
            "Currency": "USD",
            "Order Time": time,
            "Avg Price": price,
            "Executed Qty": quantity,
            "Turnover": quantity * price,
            "Order No.": "ORD-1",
        }
        sheet.append([values.get(h) for h in headers])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _payload(fund, orders, *, hk=False):
    builder = (
        build_investment_payload_from_longbridge_hk_files
        if hk
        else build_investment_payload_from_longbridge_sg_files
    )
    with patch.object(
        compat, "load_local_private_investment_evidence", return_value={}
    ):
        return builder(
            fund_details_text=fund,
            history_orders_xlsx_bytes=orders,
            fund_details_filename="H99999999 Fund Details.txt"
            if hk
            else "SG99999999 Fund Details.txt",
            history_orders_filename="Orders.xlsx",
        )


@pytest.mark.parametrize("quantity", ["0.5", "1.25"])
def test_fund_details_fractional_shares_are_not_suffix_integers(quantity):
    p = _payload(
        f"2024.Nov\nBuy Contract-Stocks\n-75.00 USD\n2024.11.15\nTQQQ.US\n{quantity} Shares\n",
        _orders(empty=True),
        hk=True,
    )
    (row,) = [r for r in p["transactions"] if r["type"] == "buy"]
    assert Decimal(row["quantity_abs"]) == Decimal(quantity)


@pytest.mark.parametrize(
    ("quantity", "price", "time", "gross"),
    [
        (1000, 12.3456, "2024-11-15 10:00:00 ET", "12345.68"),
        (10, 10, "2024-11-10 10:00:00 ET", "100.00"),
    ],
)
def test_cash_contract_and_order_cannot_create_two_trades(quantity, price, time, gross):
    fund = f"2024.Nov\nBuy Contract-Stocks\n-{gross} USD\n2024.11.15\nTQQQ.US\n{quantity} Shares\nFees for Buy Trade\n-1.50 USD\n2024.11.15\nTQQQ.US\n"
    p = _payload(fund, _orders(quantity=quantity, price=price, time=time))
    (row,) = [r for r in p["transactions"] if r["type"] == "buy"]
    assert Decimal(row["quantity_abs"]) == quantity
    assert Decimal(row["net_amount_raw"]) == -Decimal(gross) - Decimal("1.50")
    assert not any(r["type"] == "withdrawal" for r in p["transactions"])


def test_partially_cancelled_order_keeps_executed_shares_and_fee():
    fund = "2024.Nov\nBuy Contract-Stocks\n-100.00 USD\n2024.11.15\nTQQQ.US\n10 Shares\nFees for Buy Trade\n-1.50 USD\n2024.11.15\nTQQQ.US\n"
    p = _payload(fund, _orders(status="Partially Cancelled"))
    (row,) = p["transactions"]
    assert row["type"] == "buy"
    assert row["net_amount_raw"] == "-101.50"


def test_unmatched_trade_fee_is_internal_cost_not_external_withdrawal():
    row = basics._build_longbridge_cash_flow_record(
        {
            "flow_name": "Fees for Buy Trade",
            "balance": "-1.50",
            "currency": "USD",
            "time": "2024-11-15 20:00:00",
        },
        1,
        [],
    )
    assert row["type"] == "adjustment"
    assert row["normalized"]["is_cash_flow"] is False


def test_material_contract_conflict_does_not_invent_duplicate_position():
    fund = (
        "2024.Nov\nBuy Contract-Stocks\n-1000.00 USD\n2024.11.15\nTQQQ.US\n10 Shares\n"
    )
    with pytest.raises(ValueError, match="unmatched order.*overlap"):
        _payload(fund, _orders(quantity=10, price=10))


def test_ambiguous_gtc_contract_linkage_fails_closed():
    fund = "2024.Nov\nBuy Contract-Stocks\n-100.00 USD\n2024.11.15\nTQQQ.US\n10 Shares\nBuy Contract-Stocks\n-100.00 USD\n2024.11.16\nTQQQ.US\n10 Shares\n"
    with pytest.raises(ValueError, match="ambiguous matching"):
        _payload(fund, _orders(time="2024-11-10 10:00:00 ET"))


@pytest.mark.parametrize(
    ("currency", "symbol", "expected"),
    [
        ("USD", "TQQQ.US", "2024-11-16 01:00:00"),
        ("HKD", "700.HK", "2024-11-15 12:00:00"),
    ],
)
def test_file_cash_and_trade_use_same_exchange_close_utc_convention(
    currency, symbol, expected
):
    fund = f"2024.Nov\nBuy Contract-Stocks\n-100.00 {currency}\n2024.11.15\n{symbol}\n10 Shares\nCash Deposit\n1000.00 {currency}\n2024.11.15\n"
    p = _payload(fund, _orders(empty=True), hk=True)
    assert {r["datetime"] for r in p["transactions"]} == {expected}
    assert p["transactions"][0]["type"] == "deposit"


def test_bare_us_symbol_order_keeps_new_york_trade_day_without_execution_log():
    row = basics._build_longbridge_history_order_record(
        {
            "status": "Filled",
            "side": "Buy",
            "symbol": "TQQQ.US",
            "executed_quantity": "1",
            "executed_price": "10",
            "time": datetime(2026, 1, 10, 0, 30, tzinfo=timezone.utc),
        },
        1,
        [],
    )
    assert row["date"] == "2026-01-09"
    assert row["datetime"] == "2026-01-10 00:30:00"
